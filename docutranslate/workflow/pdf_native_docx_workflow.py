"""PDF -> MinerU Markdown -> native DOCX -> DOCX translation workflow."""

from __future__ import annotations

import asyncio
import os
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Self

from docx import Document as DocxDocument
from docx.oxml.ns import qn

from docutranslate.converter.md2docx import (
    NativeDocxConversionResult,
    convert_markdown_to_native_docx,
)
from docutranslate.converter.x2md.base import X2MarkdownConverter, X2MarkdownConverterConfig
from docutranslate.converter.x2md.converter_mineru import ConverterMineru
from docutranslate.converter.x2md.converter_mineru_deploy import ConverterMineruDeploy
from docutranslate.exporter.base import ExporterConfig
from docutranslate.exporter.docx.docx2docx_exporter import Docx2DocxExporter
from docutranslate.exporter.md.types import ConvertEngineType
from docutranslate.ir.document import Document
from docutranslate.workflow.base import Workflow, WorkflowConfig
from docutranslate.workflow.docx_workflow import DocxWorkflow, DocxWorkflowConfig
from docutranslate.workflow.interfaces import DocxExportable
from docutranslate.workflow.pdf_visual_inventory import (
    PdfVisualInventory,
    inspect_pdf_visuals,
)
from docutranslate.workflow.pdf_image_recovery import recover_empty_image_blocks
from docutranslate.converter.x2md.mineru_evidence import MineruMarkdownDocument


@dataclass(kw_only=True)
class PdfNativeDocxWorkflowConfig(WorkflowConfig):
    convert_engine: ConvertEngineType
    converter_config: X2MarkdownConverterConfig
    docx_workflow_config: DocxWorkflowConfig
    image_integrity_check_enabled: bool = field(default_factory=lambda: os.getenv(
        "PDF_IMAGE_INTEGRITY_CHECK_ENABLED", "false"
    ).strip().lower() in {"1", "true", "yes", "on"})


class PdfNativeDocxWorkflow(
    Workflow[PdfNativeDocxWorkflowConfig, Document, Document],
    DocxExportable[ExporterConfig],
):
    _converter_factory = {
        "mineru": ConverterMineru,
        "mineru_deploy": ConverterMineruDeploy,
    }
    _max_parse_attempts = 2
    _min_substantive_image_bytes = 1024
    _required_table_borders = frozenset(
        {"top", "left", "bottom", "right", "insideH", "insideV"}
    )

    def __init__(self, config: PdfNativeDocxWorkflowConfig):
        super().__init__(config=config)
        if config.convert_engine not in self._converter_factory:
            raise ValueError("pdf_native_docx only supports mineru and mineru_deploy")
        self._docx_workflow: DocxWorkflow | None = None
        self._conversion_diagnostics: dict = {}

    def _prepare(self) -> tuple[X2MarkdownConverter, Document]:
        if self.document_original is None:
            raise RuntimeError("File has not been read yet. Call read_path or read_bytes first.")
        if self.document_original.suffix.lower() != ".pdf":
            raise ValueError("pdf_native_docx only supports PDF input")
        converter_class = self._converter_factory[self.config.convert_engine]
        converter = converter_class(self.config.converter_config)
        return converter, self.document_original

    def _translate_docx(self, document: Document) -> DocxWorkflow:
        workflow = DocxWorkflow(self.config.docx_workflow_config)
        workflow.read_bytes(document.content, stem=document.stem or "document", suffix=".docx")
        workflow.translate()
        return workflow

    async def _translate_docx_async(self, document: Document) -> DocxWorkflow:
        workflow = DocxWorkflow(self.config.docx_workflow_config)
        workflow.read_bytes(document.content, stem=document.stem or "document", suffix=".docx")
        await workflow.translate_async()
        return workflow

    def _accept_docx_workflow(self, workflow: DocxWorkflow) -> None:
        if workflow.document_translated is None:
            raise RuntimeError("DOCX translation did not produce a document")
        self._validate_translated_docx(workflow.document_translated.content)
        self._docx_workflow = workflow
        self.document_translated = workflow.document_translated

    def _inspect_source_pdf_visuals(self, source: Document) -> PdfVisualInventory | None:
        try:
            return inspect_pdf_visuals(
                source.content,
                min_substantive_image_bytes=self._min_substantive_image_bytes,
            )
        except Exception as exc:
            self.logger.warning("Could not inspect source PDF visuals: %r", exc)
            return None

    @staticmethod
    def _validate_scan_page_content(
        inventory: PdfVisualInventory | None,
        markdown: Document,
    ) -> list[int]:
        if inventory is None or not inventory.scan_pages:
            return []
        content_pages = getattr(markdown, "mineru_content_pages", None)
        if content_pages is None:
            raise RuntimeError(
                "PDF_SCAN_PAGE_EVIDENCE_MISSING: MinerU did not return page-level "
                f"content evidence for scan pages={','.join(map(str, inventory.scan_pages))}"
            )
        covered_pages = sorted(
            page_number
            for page_number in inventory.scan_pages
            if page_number - 1 in content_pages
        )
        missing_pages = sorted(set(inventory.scan_pages) - set(covered_pages))
        if missing_pages:
            raise RuntimeError(
                "PDF_SCAN_PAGE_CONTENT_MISSING: MinerU returned no translatable content "
                f"for scan pages={','.join(map(str, missing_pages))}"
            )
        return covered_pages

    @staticmethod
    def _candidate_image_count(candidate: NativeDocxConversionResult) -> int:
        return int(candidate.diagnostics["validation"]["docx_inline_images"])

    @classmethod
    def _candidate_score(cls, candidate: NativeDocxConversionResult) -> tuple[int, int, float]:
        validation = candidate.diagnostics["validation"]
        return (
            cls._candidate_image_count(candidate),
            int(validation["docx_tables"]),
            float(validation["ast_to_docx_character_coverage"]),
        )

    def _select_conversion_candidate(
        self,
        candidates: list[NativeDocxConversionResult],
        inventory: PdfVisualInventory | None,
        scan_pages_with_content: list[int],
    ) -> NativeDocxConversionResult:
        selected_index, selected = max(
            enumerate(candidates), key=lambda item: self._candidate_score(item[1])
        )
        candidate_images = [self._candidate_image_count(item) for item in candidates]
        source_diagnostics = (
            inventory.diagnostics()
            if inventory is not None
            else {
                "source_pdf_image_objects": None,
                "source_pdf_substantive_images": None,
                "source_pdf_inline_images": None,
                "source_pdf_scan_images": None,
                "source_pdf_scan_pages": [],
                "source_pdf_ignored_image_masks": None,
                "source_pdf_ignored_small_images": None,
            }
        )
        diagnostics = {
            **selected.diagnostics,
            "image_integrity_check_enabled": self.config.image_integrity_check_enabled,
            **source_diagnostics,
            "scan_pages_with_mineru_content": scan_pages_with_content,
            "mineru_parse_attempts": len(candidates),
            "selected_parse_attempt": selected_index + 1,
            "candidate_inline_images": candidate_images,
        }
        selected = NativeDocxConversionResult(
            document=selected.document,
            diagnostics=diagnostics,
        )
        selected_images = self._candidate_image_count(selected)
        source_inline_images = inventory.inline_images if inventory is not None else None
        if (self.config.image_integrity_check_enabled
                and source_inline_images is not None and selected_images < source_inline_images):
            raise RuntimeError(
                "Conservative PDF inline-image validation failed after "
                f"{len(candidates)} MinerU attempts: "
                f"source_inline_images={source_inline_images}, "
                f"docx_inline_images={selected_images}"
            )
        return selected

    def _validate_translated_docx(self, content: bytes) -> None:
        doc = DocxDocument(BytesIO(content))
        expected = self._conversion_diagnostics["validation"]
        bordered_tables = 0
        for table in doc.tables:
            borders = table._tbl.tblPr.find(qn("w:tblBorders"))
            border_names = (
                {child.tag.rsplit("}", 1)[-1] for child in borders}
                if borders is not None
                else set()
            )
            if self._required_table_borders <= border_names:
                bordered_tables += 1
        validation = {
            "expected_tables": int(expected["docx_tables"]),
            "docx_tables": len(doc.tables),
            "tables_with_all_borders": bordered_tables,
            "expected_images": int(expected["docx_inline_images"]),
            "docx_inline_images": len(doc.inline_shapes),
        }
        if (
            validation["docx_tables"] != validation["expected_tables"]
            or validation["tables_with_all_borders"] != validation["docx_tables"]
            or (self.config.image_integrity_check_enabled
                and validation["docx_inline_images"] != validation["expected_images"])
        ):
            raise RuntimeError(f"Translated DOCX structural validation failed: {validation}")
        self._conversion_diagnostics["post_translation_validation"] = validation

    def _convert_native_candidate(self, source, markdown, converter):
        recovered = 0
        # Deployment archives use source page indexes only for a full-page start.
        # Cloud chunk archives require separate offset handling; keep their gate.
        if (isinstance(converter, ConverterMineruDeploy)
                and converter.config.start_page_id == 0 and converter.attachments):
            try:
                repaired, recovered = recover_empty_image_blocks(
                    source.content, markdown.content.decode('utf-8'),
                    converter.attachments[-1].document.content,
                )
            except Exception:
                self.logger.warning('Source image recovery unavailable; retaining strict validation')
            if recovered:
                markdown = MineruMarkdownDocument.from_bytes(
                    repaired.encode('utf-8'), suffix='.md', stem=markdown.stem,
                    mineru_content_pages=getattr(markdown, 'mineru_content_pages', None),
                )
                self.logger.info('Restored %s source images from empty MinerU blocks', recovered)
        converted = convert_markdown_to_native_docx(markdown, logger=self.logger)
        return NativeDocxConversionResult(
            document=converted.document,
            diagnostics={**converted.diagnostics, 'recovered_source_images': recovered},
        )

    def translate(self) -> Self:
        converter, source = self._prepare()
        inventory = self._inspect_source_pdf_visuals(source)
        source_inline_images = inventory.inline_images if inventory is not None else None
        candidates = []
        scan_pages_with_content: list[int] = []
        for _ in range(self._max_parse_attempts):
            self.progress_tracker.update(percent=10, message="正在解析文档...")
            markdown = converter.convert(source)
            scan_pages_with_content = self._validate_scan_page_content(inventory, markdown)
            self.progress_tracker.update(percent=20, message="正在生成 Word...")
            candidates.append(self._convert_native_candidate(source, markdown, converter))
            if (
                not self.config.image_integrity_check_enabled
                or source_inline_images is None
                or self._candidate_image_count(candidates[-1]) >= source_inline_images
            ):
                break
            self.logger.warning(
                "MinerU output is missing source PDF inline images; retrying parse "
                "(%s/%s images)",
                self._candidate_image_count(candidates[-1]),
                source_inline_images,
            )
        converted = self._select_conversion_candidate(
            candidates,
            inventory,
            scan_pages_with_content,
        )
        self._conversion_diagnostics = converted.diagnostics
        self._accept_docx_workflow(self._translate_docx(converted.document))
        return self

    async def translate_async(self) -> Self:
        converter, source = self._prepare()
        inventory = await asyncio.to_thread(self._inspect_source_pdf_visuals, source)
        source_inline_images = inventory.inline_images if inventory is not None else None
        candidates = []
        scan_pages_with_content: list[int] = []
        for _ in range(self._max_parse_attempts):
            self.progress_tracker.update(percent=10, message="正在解析文档...")
            markdown = await converter.convert_async(source)
            scan_pages_with_content = self._validate_scan_page_content(inventory, markdown)
            self.progress_tracker.update(percent=20, message="正在生成 Word...")
            candidates.append(
                await asyncio.to_thread(
                    self._convert_native_candidate,
                    source,
                    markdown,
                    converter,
                )
            )
            if (
                not self.config.image_integrity_check_enabled
                or source_inline_images is None
                or self._candidate_image_count(candidates[-1]) >= source_inline_images
            ):
                break
            self.logger.warning(
                "MinerU output is missing source PDF inline images; retrying parse "
                "(%s/%s images)",
                self._candidate_image_count(candidates[-1]),
                source_inline_images,
            )
        converted = self._select_conversion_candidate(
            candidates,
            inventory,
            scan_pages_with_content,
        )
        self._conversion_diagnostics = converted.diagnostics
        self._accept_docx_workflow(await self._translate_docx_async(converted.document))
        return self

    def get_statistics(self) -> dict:
        statistics = self._docx_workflow.get_statistics() if self._docx_workflow else {}
        return {**statistics, "conversion": self._conversion_diagnostics}

    def export_to_docx(self, _: ExporterConfig | None = None) -> bytes:
        return self._export(Docx2DocxExporter()).content

    def save_as_docx(
        self,
        name: str | None = None,
        output_dir: Path | str = "./output",
        _: ExporterConfig | None = None,
    ) -> Self:
        self._save(exporter=Docx2DocxExporter(), name=name, output_dir=output_dir)
        return self
