from __future__ import annotations

from base64 import b64encode
from io import BytesIO
import logging
from random import Random
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from PIL import Image
from pydantic import TypeAdapter, ValidationError
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    BooleanObject,
    DecodedStreamObject,
    DictionaryObject,
    EncodedStreamObject,
    NameObject,
    NumberObject,
)
import pytest

from docutranslate.converter.md2docx import convert_markdown_to_native_docx
from docutranslate.converter.x2md.mineru_evidence import (
    MineruMarkdownDocument,
    mineru_content_pages_from_zip,
)
from docutranslate.core.factory import create_workflow_from_payload
from docutranslate.core.schemas import PdfNativeDocxWorkflowParams, TranslatePayload
from docutranslate.ir.markdown_document import MarkdownDocument
from docutranslate.utils.markdown_utils import (
    MaskDict,
    embed_inline_image_from_zip,
    placeholder2uris,
    uris2placeholder,
)
from docutranslate.workflow.pdf_native_docx_workflow import PdfNativeDocxWorkflow


def _payload(**overrides):
    values = {
        "workflow_type": "pdf_native_docx",
        "skip_translate": True,
        "convert_engine": "mineru_deploy",
        "mineru_deploy_base_url": "http://mineru:8000",
        "translation_review_enable": True,
    }
    values.update(overrides)
    return TypeAdapter(TranslatePayload).validate_python(values)


def _pdf_with_raster_images(count: int) -> bytes:
    random = Random(42)
    images = [
        Image.frombytes(
            "RGB",
            (200, 200),
            bytes(random.randrange(256) for _ in range(200 * 200 * 3)),
        )
        for _ in range(count)
    ]
    output = BytesIO()
    images[0].save(output, format="PDF", save_all=True, append_images=images[1:])
    return output.getvalue()


def _pdf_with_inline_raster_images(count: int) -> bytes:
    source = PdfReader(BytesIO(_pdf_with_raster_images(count)))
    output = BytesIO()
    writer = PdfWriter()
    for source_page in source.pages:
        page = writer.add_page(source_page)
        page.mediabox.upper_right = (1000, 1000)
        page.cropbox.upper_right = (1000, 1000)
    writer.write(output)
    return output.getvalue()


def _pdf_with_full_page_scans_and_image_mask() -> bytes:
    random = Random(73)
    writer = PdfWriter()

    def add_image_xobject(*, width: int, height: int, image_mask: bool = False):
        if image_mask:
            stream = DecodedStreamObject()
            row_bytes = (width + 7) // 8
            stream.set_data(random.randbytes(row_bytes * height))
            stream.update(
                {
                    NameObject("/Type"): NameObject("/XObject"),
                    NameObject("/Subtype"): NameObject("/Image"),
                    NameObject("/Width"): NumberObject(width),
                    NameObject("/Height"): NumberObject(height),
                    NameObject("/ImageMask"): BooleanObject(True),
                    NameObject("/BitsPerComponent"): NumberObject(1),
                }
            )
        else:
            encoded = BytesIO()
            pixels = random.randbytes(width * height * 3)
            Image.frombytes("RGB", (width, height), pixels).save(
                encoded,
                format="JPEG",
                quality=80,
            )
            stream = EncodedStreamObject()
            stream._data = encoded.getvalue()
            stream.update(
                {
                    NameObject("/Type"): NameObject("/XObject"),
                    NameObject("/Subtype"): NameObject("/Image"),
                    NameObject("/Width"): NumberObject(width),
                    NameObject("/Height"): NumberObject(height),
                    NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                    NameObject("/BitsPerComponent"): NumberObject(8),
                    NameObject("/Filter"): NameObject("/DCTDecode"),
                }
            )
        return writer._add_object(stream)

    first_page = writer.add_blank_page(width=612, height=792)
    first_scan = add_image_xobject(width=612, height=792)
    first_mask = add_image_xobject(width=200, height=240, image_mask=True)
    first_page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/XObject"): DictionaryObject(
                {
                    NameObject("/ScanPage1"): first_scan,
                    NameObject("/Stencil"): first_mask,
                }
            )
        }
    )
    first_content = DecodedStreamObject()
    first_content.set_data(
        b"q 612 0 0 792 0 0 cm /ScanPage1 Do Q "
        b"q 440 0 0 630 86 68 cm /Stencil Do Q"
    )
    first_page[NameObject("/Contents")] = writer._add_object(first_content)

    second_page = writer.add_blank_page(width=595, height=842)
    second_scan = add_image_xobject(width=595, height=842)
    second_page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/XObject"): DictionaryObject(
                {NameObject("/ScanPage2"): second_scan}
            )
        }
    )
    second_content = DecodedStreamObject()
    second_content.set_data(b"q 595 0 0 842 0 0 cm /ScanPage2 Do Q")
    second_page[NameObject("/Contents")] = writer._add_object(second_content)

    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_native_docx_schema_accepts_only_mineru_engines():
    payload = _payload(mineru_deploy_backend="hybrid-engine")
    assert isinstance(payload, PdfNativeDocxWorkflowParams)
    assert payload.mineru_deploy_backend == "hybrid-engine"

    with pytest.raises(ValidationError):
        _payload(convert_engine="pypdf")


def test_factory_reuses_docx_translation_options():
    workflow = create_workflow_from_payload(_payload())

    assert isinstance(workflow, PdfNativeDocxWorkflow)
    translator = workflow.config.docx_workflow_config.translator_config
    assert translator.skip_translate is True
    assert translator.translation_review_enable is True
    assert workflow.config.convert_engine == "mineru_deploy"
    assert workflow.config.converter_config.return_content_list is True


def test_mineru_content_list_reports_only_pages_with_translatable_content():
    mineru_zip = BytesIO()
    with ZipFile(mineru_zip, "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "result/source_content_list.json",
            """[
                {"type":"text","text":"page zero","page_idx":0},
                {"type":"image","img_path":"images/only.png","page_idx":1},
                {"type":"table","table_body":"<table><tr><td>x</td></tr></table>","page_idx":2},
                {"type":"page_number","text":"4","page_idx":3}
            ]""",
        )

    assert mineru_content_pages_from_zip(mineru_zip.getvalue()) == frozenset({0, 2})


def test_native_conversion_preserves_mineru_html_table_and_math():
    markdown = MarkdownDocument.from_bytes(
        content=(
            b"Before\n\n"
            b"<table><tr><td>METRIC</td><td>$<E_{o}$</td></tr>"
            b"<tr><td>HMS</td><td>$\\geq 9.3$</td></tr></table>\n\n"
            b"After\n"
        ),
        suffix=".md",
        stem="sample",
    )

    result = convert_markdown_to_native_docx(markdown)
    docx = DocxDocument(BytesIO(result.document.content))

    assert result.document.suffix == ".docx"
    assert len(docx.tables) == 1
    assert "<E" in "".join(cell.text for row in docx.tables[0].rows for cell in row.cells)
    assert result.diagnostics["converted_html_tables"] == 1
    assert result.diagnostics["validation"]["ast_to_docx_character_coverage"] == 1.0
    borders = docx.tables[0]._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None
    assert {
        child.tag.rsplit("}", 1)[-1]: child.get(qn("w:val"))
        for child in borders
    } == {
        "top": "single",
        "left": "single",
        "bottom": "single",
        "right": "single",
        "insideH": "single",
        "insideV": "single",
    }
    assert all(
        row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
        for row in docx.tables[0].rows
    )
    assert all(
        row._tr.get_or_add_trPr().find(qn("w:tblHeader")) is None
        for row in docx.tables[0].rows
    )


def test_native_conversion_preserves_inline_images():
    image = BytesIO()
    Image.new("RGB", (2, 2), "red").save(image, format="PNG")
    markdown = MarkdownDocument.from_bytes(
        content=(
            "Before\n\n"
            f"![sample](data:image/png;base64,{b64encode(image.getvalue()).decode()})\n\n"
            "After\n"
        ).encode(),
        suffix=".md",
        stem="image-sample",
    )

    result = convert_markdown_to_native_docx(markdown)
    docx = DocxDocument(BytesIO(result.document.content))

    assert result.diagnostics["validation"]["expected_images"] == 1
    assert len(docx.inline_shapes) == 1
    assert result.diagnostics["validation"]["docx_inline_images"] == 1


def test_native_conversion_preserves_html_table_images_from_mineru_zip():
    outside_image = BytesIO()
    Image.new("RGB", (2, 2), "red").save(outside_image, format="PNG")
    table_image = BytesIO()
    Image.new("RGB", (2, 2), "blue").save(table_image, format="PNG")
    mineru_zip = BytesIO()
    with ZipFile(mineru_zip, "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "result/document.md",
            "Before\n\n"
            "![outside](images/outside.png)\n\n"
            '<table><tr><td><img src="images/table.png" alt="inside table"></td></tr></table>\n\n'
            "After\n",
        )
        archive.writestr("result/images/outside.png", outside_image.getvalue())
        archive.writestr("result/images/table.png", table_image.getvalue())

    embedded = embed_inline_image_from_zip(mineru_zip.getvalue())
    assert embedded is not None
    assert embedded.count("data:image/png;base64,") == 2

    markdown = MarkdownDocument.from_bytes(
        content=embedded.encode(),
        suffix=".md",
        stem="html-table-images",
    )
    result = convert_markdown_to_native_docx(markdown)
    docx = DocxDocument(BytesIO(result.document.content))

    assert result.diagnostics["validation"]["expected_images"] == 2
    assert result.diagnostics["validation"]["docx_inline_images"] == 2
    assert len(docx.inline_shapes) == 2


def test_markdown_translation_masks_and_restores_html_images():
    source = (
        'Before\n\n<table><tr><td><img src="data:image/png;base64,AAAA" '
        'alt="inside table"></td></tr></table>\n\nAfter\n'
    )
    mask = MaskDict()

    masked = uris2placeholder(source, mask)

    assert "data:image/png;base64" not in masked
    assert placeholder2uris(masked, mask) == source


@pytest.mark.asyncio
async def test_composite_workflow_returns_only_translated_docx(monkeypatch):
    class FakeMineruConverter:
        def __init__(self, _config):
            pass

        async def convert_async(self, document):
            return MarkdownDocument.from_bytes(
                content=(
                    b"# Heading\n\nBody text for translation.\n\n"
                    b"<table><tr><td>A</td><td>B</td></tr>"
                    b"<tr><td>C</td><td>D</td></tr></table>\n"
                ),
                suffix=".md",
                stem=document.stem,
            )

    payload = _payload()
    workflow = create_workflow_from_payload(payload, logger=logging.getLogger("test-native-docx"))
    monkeypatch.setitem(workflow._converter_factory, "mineru_deploy", FakeMineruConverter)
    workflow.read_bytes(b"%PDF-1.7\n", stem="sample", suffix=".pdf")

    await workflow.translate_async()

    exported = workflow.export_to_docx()
    result_docx = DocxDocument(BytesIO(exported))
    assert "Heading" in "\n".join(paragraph.text for paragraph in result_docx.paragraphs)
    assert len(result_docx.tables) == 1
    borders = result_docx.tables[0]._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None
    assert {child.tag.rsplit("}", 1)[-1] for child in borders} == {
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    }
    assert workflow.get_attachment().attachment_dict == {}
    assert workflow.get_statistics()["conversion"]["validation"]["ast_to_docx_order_coverage"] == 1.0
    assert workflow.get_statistics()["conversion"]["post_translation_validation"][
        "tables_with_all_borders"
    ] == 1


@pytest.mark.asyncio
@pytest.mark.parametrize("asynchronous", [False, True])
@pytest.mark.parametrize("enabled", [False, True])
async def test_pdf_image_check_switch_defaults_off(monkeypatch, asynchronous, enabled):
    monkeypatch.delenv("PDF_IMAGE_INTEGRITY_CHECK_ENABLED", raising=False)
    if enabled:
        monkeypatch.setenv("PDF_IMAGE_INTEGRITY_CHECK_ENABLED", "true")

    class MissingImagesConverter:
        calls = 0

        def __init__(self, _config):
            pass

        def convert(self, document):
            type(self).calls += 1
            return MarkdownDocument.from_bytes(b"Recognized text", suffix=".md", stem=document.stem)

        async def convert_async(self, document):
            return self.convert(document)

    workflow = create_workflow_from_payload(_payload())
    monkeypatch.setitem(workflow._converter_factory, "mineru_deploy", MissingImagesConverter)
    workflow.read_bytes(_pdf_with_inline_raster_images(1), stem="missing", suffix=".pdf")

    async def run():
        if asynchronous:
            await workflow.translate_async()
        else:
            workflow.translate()

    if enabled:
        with pytest.raises(RuntimeError, match="source_inline_images=1, docx_inline_images=0"):
            await run()
        assert MissingImagesConverter.calls == 2
    else:
        await run()
        assert MissingImagesConverter.calls == 1
        assert len(DocxDocument(BytesIO(workflow.export_to_docx())).inline_shapes) == 0
        assert workflow.get_statistics()["conversion"]["image_integrity_check_enabled"] is False


@pytest.mark.parametrize("enabled", [False, True])
def test_pdf_image_switch_controls_post_translation_image_check(monkeypatch, enabled):
    monkeypatch.setenv("PDF_IMAGE_INTEGRITY_CHECK_ENABLED", str(enabled))
    workflow = create_workflow_from_payload(_payload())
    workflow._conversion_diagnostics = {"validation": {"docx_tables": 0, "docx_inline_images": 1}}
    content = BytesIO()
    DocxDocument().save(content)
    if enabled:
        with pytest.raises(RuntimeError, match="Translated DOCX structural validation failed"):
            workflow._validate_translated_docx(content.getvalue())
    else:
        workflow._validate_translated_docx(content.getvalue())
        workflow._conversion_diagnostics["validation"]["docx_tables"] = 1
        with pytest.raises(RuntimeError, match="Translated DOCX structural validation failed"):
            workflow._validate_translated_docx(content.getvalue())


@pytest.mark.asyncio
async def test_native_workflow_retries_mineru_when_source_images_are_missing(monkeypatch):
    monkeypatch.setenv("PDF_IMAGE_INTEGRITY_CHECK_ENABLED", "true")
    encoded_images = []
    for color in ("red", "blue"):
        image = BytesIO()
        Image.new("RGB", (80, 80), color).save(image, format="PNG")
        encoded_images.append(b64encode(image.getvalue()).decode())

    class FlakyMineruConverter:
        calls = 0

        def __init__(self, _config):
            pass

        async def convert_async(self, document):
            self.__class__.calls += 1
            image_count = self.__class__.calls
            markdown = "\n\n".join(
                f"![figure-{index}](data:image/png;base64,{encoded_images[index]})"
                for index in range(image_count)
            )
            return MarkdownDocument.from_bytes(
                content=markdown.encode(),
                suffix=".md",
                stem=document.stem,
            )

    workflow = create_workflow_from_payload(
        _payload(), logger=logging.getLogger("test-native-docx-image-retry")
    )
    monkeypatch.setitem(workflow._converter_factory, "mineru_deploy", FlakyMineruConverter)
    workflow.read_bytes(
        _pdf_with_inline_raster_images(2), stem="flaky-images", suffix=".pdf"
    )

    await workflow.translate_async()

    exported = DocxDocument(BytesIO(workflow.export_to_docx()))
    diagnostics = workflow.get_statistics()["conversion"]
    assert FlakyMineruConverter.calls == 2
    assert len(exported.inline_shapes) == 2
    assert diagnostics["source_pdf_substantive_images"] == 2
    assert diagnostics["mineru_parse_attempts"] == 2
    assert diagnostics["selected_parse_attempt"] == 2
    assert diagnostics["candidate_inline_images"] == [1, 2]
    assert diagnostics["post_translation_validation"]["docx_inline_images"] == 2


@pytest.mark.asyncio
async def test_native_workflow_accepts_ocr_covered_scan_pages_without_inline_images(monkeypatch):
    class ScanAwareMineruConverter:
        calls = 0

        def __init__(self, _config):
            pass

        async def convert_async(self, document):
            self.__class__.calls += 1
            markdown = MineruMarkdownDocument.from_bytes(
                content=b"# OCR page one\n\nOCR page two\n",
                suffix=".md",
                stem=document.stem,
                mineru_content_pages=frozenset({0, 1}),
            )
            return markdown

    workflow = create_workflow_from_payload(
        _payload(), logger=logging.getLogger("test-native-docx-scan-pages")
    )
    monkeypatch.setitem(
        workflow._converter_factory,
        "mineru_deploy",
        ScanAwareMineruConverter,
    )
    workflow.read_bytes(
        _pdf_with_full_page_scans_and_image_mask(),
        stem="scanned-pages",
        suffix=".pdf",
    )

    await workflow.translate_async()

    diagnostics = workflow.get_statistics()["conversion"]
    assert ScanAwareMineruConverter.calls == 1
    assert diagnostics["source_pdf_image_objects"] == 3
    assert diagnostics["source_pdf_ignored_image_masks"] == 1
    assert diagnostics["source_pdf_scan_images"] == 2
    assert diagnostics["source_pdf_scan_pages"] == [1, 2]
    assert diagnostics["source_pdf_inline_images"] == 0
    assert diagnostics["scan_pages_with_mineru_content"] == [1, 2]
    assert diagnostics["candidate_inline_images"] == [0]


@pytest.mark.asyncio
async def test_native_workflow_rejects_scan_page_without_ocr_content(monkeypatch):
    class IncompleteScanMineruConverter:
        def __init__(self, _config):
            pass

        async def convert_async(self, document):
            markdown = MineruMarkdownDocument.from_bytes(
                content=b"# OCR page one only\n",
                suffix=".md",
                stem=document.stem,
                mineru_content_pages=frozenset({0}),
            )
            return markdown

    workflow = create_workflow_from_payload(
        _payload(), logger=logging.getLogger("test-native-docx-scan-page-gap")
    )
    monkeypatch.setitem(
        workflow._converter_factory,
        "mineru_deploy",
        IncompleteScanMineruConverter,
    )
    workflow.read_bytes(
        _pdf_with_full_page_scans_and_image_mask(),
        stem="scanned-page-gap",
        suffix=".pdf",
    )

    with pytest.raises(
        RuntimeError,
        match=r"PDF_SCAN_PAGE_CONTENT_MISSING.*pages=2",
    ):
        await workflow.translate_async()


@pytest.mark.asyncio
async def test_native_workflow_rejects_scan_pages_without_page_evidence(monkeypatch):
    class EvidenceFreeMineruConverter:
        def __init__(self, _config):
            pass

        async def convert_async(self, document):
            return MarkdownDocument.from_bytes(
                content=b"# Unscoped OCR output\n",
                suffix=".md",
                stem=document.stem,
            )

    workflow = create_workflow_from_payload(
        _payload(), logger=logging.getLogger("test-native-docx-scan-evidence-gap")
    )
    monkeypatch.setitem(
        workflow._converter_factory,
        "mineru_deploy",
        EvidenceFreeMineruConverter,
    )
    workflow.read_bytes(
        _pdf_with_full_page_scans_and_image_mask(),
        stem="scanned-evidence-gap",
        suffix=".pdf",
    )

    with pytest.raises(
        RuntimeError,
        match=r"PDF_SCAN_PAGE_EVIDENCE_MISSING.*pages=1,2",
    ):
        await workflow.translate_async()
